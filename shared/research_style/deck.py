"""Build editable PPTX slides from a JSON outline and shared style."""
import argparse
from io import BytesIO
import math
from pathlib import Path
import unicodedata

from .core import (artifact_manifest, atomic_bytes, load_style, manifest_path,
                   preflight_outputs, read_json, write_json)


def text_units(text):
    return sum(1 if unicodedata.east_asian_width(c) in "WF" else .55 for c in text)


def text_lines(text, width_mm, font_pt):
    units_per_line = width_mm * 72 / 25.4 / font_pt
    return sum(max(1, math.ceil(text_units(line) / units_per_line)) for line in text.split("\n"))


def check_deck(spec):
    if not isinstance(spec, dict) or not isinstance(spec.get("slides"), list) or not spec["slides"]:
        raise ValueError("Deck JSON requires a nonempty slides array")
    for slide in spec["slides"]:
        if not isinstance(slide, dict) or not isinstance(slide.get("title"), str) or not slide["title"].strip():
            raise ValueError("Every slide requires a nonempty title")
        if not isinstance(slide.get("bullets", []), list) or not all(isinstance(x, str) and x.strip() for x in slide.get("bullets", [])):
            raise ValueError("Slide bullets must be nonempty strings")
        for key in ("image", "caption"):
            if key in slide and not isinstance(slide[key], str):
                raise ValueError(f"Slide {key} must be a string")


def build(spec, source_dir, style):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Mm, Pt
    from PIL import Image
    check_deck(spec)
    cfg, palette, fonts = style["slides"], style["palette"], style["typography"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(cfg["width_mm"]), Mm(cfg["height_mm"])
    margin = cfg["margin_mm"]
    width = cfg["width_mm"] - 2 * margin
    title_height = cfg["title_pt"] * 25.4 / 72 * 2.6
    body_top = margin + title_height + 4
    body_height = cfg["height_mm"] - margin - body_top - 16
    if body_height < cfg["body_pt"] * 25.4 / 72 * 2:
        raise ValueError("Slide height/font/margins leave insufficient content space")

    def fill_text(shape, text, pt, color, bold=False):
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.15
            paragraph.alignment = PP_ALIGN.LEFT
            # Explicitly disable inherited placeholder bullets; the input string
            # already contains the intended bullet glyph. Empty spacing rows
            # must not acquire a bullet from the PowerPoint master.
            props = paragraph._p.get_or_add_pPr()
            props.set("marL", "0")
            props.set("indent", "0")
            props.append(OxmlElement("a:buNone"))
            paragraph.font.name, paragraph.font.size = fonts["latin"], Pt(pt)
            run = paragraph.add_run()
            run.text = line
            run.font.name, run.font.size, run.font.bold = fonts["latin"], Pt(pt), bold
            run.font.color.rgb = RGBColor.from_string(color[1:])
            ea = OxmlElement("a:ea")
            ea.set("typeface", fonts["cjk"])
            run._r.get_or_add_rPr().append(ea)
        return frame

    for entry in spec["slides"]:
        if text_lines(entry["title"], width, cfg["title_pt"]) > 2:
            raise ValueError(f"Title exceeds two lines at configured size: {entry['title']}")
        image = (source_dir / entry["image"]).resolve() if entry.get("image") else None
        if image and (not image.is_file() or image.suffix.lower() not in (".png", ".jpg", ".jpeg")):
            raise ValueError(f"Slide image must be an existing PNG/JPEG: {image}")
        content_width = width * .43 if image and entry.get("bullets") else width
        capacity = int(body_height / (cfg["body_pt"] * 25.4 / 72 * 1.4))
        pages, current, used = [], [], 0
        for bullet in entry.get("bullets", []):
            needed = text_lines("• " + bullet, content_width, cfg["body_pt"]) + 1
            if needed > capacity:
                raise ValueError("One bullet exceeds slide height; split it into smaller statements")
            if current and used + needed > capacity:
                pages.append(current)
                current, used = [], 0
            current.append(bullet)
            used += needed
        pages.append(current)
        caption = entry.get("caption", "")
        if text_lines(caption, width - 10, cfg["caption_pt"]) > 2:
            raise ValueError("Caption exceeds two lines; shorten it or move detail into the body")
        for page_index, bullets in enumerate(pages):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(palette["background"][1:])
            title = slide.shapes.title
            title.left, title.top, title.width, title.height = Mm(margin), Mm(margin), Mm(width), Mm(title_height)
            fill_text(title, entry["title"], cfg["title_pt"], palette["accent"], True)
            body = slide.placeholders[1]
            body.left, body.top, body.width, body.height = Mm(margin), Mm(body_top), Mm(content_width), Mm(body_height)
            fill_text(body, "\n\n".join("• " + b for b in bullets), cfg["body_pt"], palette["text"])
            if image:
                box_x = margin + width * .48 if bullets else margin
                box_w = width * .52 if bullets else width
                with Image.open(image) as bitmap:
                    ratio = bitmap.width / bitmap.height
                pic_w = min(box_w, body_height * ratio)
                pic_h = pic_w / ratio
                slide.shapes.add_picture(str(image), Mm(box_x + (box_w - pic_w) / 2),
                                         Mm(body_top + (body_height - pic_h) / 2), Mm(pic_w), Mm(pic_h))
            footer = slide.shapes.add_textbox(Mm(margin), Mm(cfg["height_mm"] - margin - 11), Mm(width - 10), Mm(11))
            continuation = f" ({page_index + 1}/{len(pages)})" if len(pages) > 1 else ""
            fill_text(footer, caption + continuation, cfg["caption_pt"], palette["muted"])
            number = slide.shapes.add_textbox(Mm(cfg["width_mm"] - margin - 8), Mm(cfg["height_mm"] - margin - 8), Mm(8), Mm(8))
            fill_text(number, str(len(prs.slides)), cfg["caption_pt"], palette["muted"])
    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue(), len(prs.slides)


def main(skill_dir):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("outline", type=Path)
    p.add_argument("--style")
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--force", action="store_true")
    a = p.parse_args()
    if a.output.suffix.lower() != ".pptx":
        raise ValueError("Output must end in .pptx")
    style, source = load_style(skill_dir, a.style)
    spec = read_json(a.outline)
    check_deck(spec)
    images = [(a.outline.parent / s["image"]) for s in spec["slides"] if s.get("image")]
    preflight_outputs([a.output, manifest_path(a.output)], a.force, [source, a.outline, *images])
    content, count = build(spec, a.outline.parent, style)
    atomic_bytes(a.output, content)
    write_json(manifest_path(a.output), artifact_manifest(style, source, "python-pptx", slides=count,
               warnings=["Text fit uses conservative estimates; render in PowerPoint/LibreOffice for final visual QA. Fonts must be installed on the viewing machine."]))
    print(f"{a.output} ({count} slides)")
