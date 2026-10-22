"""Assemble a DOCX paper from supplied content and explicit citation keys."""
import argparse
from io import BytesIO
from pathlib import Path
import re

from .core import (artifact_manifest, atomic_bytes, load_style, manifest_path,
                   preflight_outputs, read_json, write_json)

CITATION = re.compile(r"\[@([a-zA-Z0-9_.:-]+)\]")


def check_paper(spec):
    if not isinstance(spec, dict) or not isinstance(spec.get("title"), str) or not spec["title"].strip():
        raise ValueError("Paper requires a nonempty title")
    for key in ("authors", "abstract"):
        if key in spec and not isinstance(spec[key], str):
            raise ValueError(f"{key} must be a string")
    if not isinstance(spec.get("sections"), list) or not spec["sections"]:
        raise ValueError("Paper requires a nonempty sections array")
    for section in spec["sections"]:
        if not isinstance(section, dict) or not isinstance(section.get("heading"), str):
            raise ValueError("Every section requires a heading")
        if not isinstance(section.get("paragraphs"), list) or not all(isinstance(v, str) for v in section["paragraphs"]):
            raise ValueError("Section paragraphs must be a string array")
    refs = spec.get("references", [])
    if not isinstance(refs, list):
        raise ValueError("References must be an array")
    index = {}
    for ref in refs:
        if not isinstance(ref, dict) or not all(isinstance(ref.get(k), str) and ref[k].strip() for k in ("id", "authors", "year", "title", "source")):
            raise ValueError("Each reference requires nonempty string id/authors/year/title/source")
        if ref["id"] in index or not re.fullmatch(r"[a-zA-Z0-9_.:-]+", ref["id"]):
            raise ValueError(f"Invalid or duplicate reference ID: {ref['id']}")
        index[ref["id"]] = ref
    texts = [spec.get("abstract", ""), *(t for section in spec["sections"] for t in section["paragraphs"])]
    cited = list(dict.fromkeys(key for text in texts for key in CITATION.findall(text)))
    missing = set(cited) - set(index)
    if missing:
        raise ValueError(f"Unresolved citations: {sorted(missing)}; supply sources, do not invent them")
    return index, cited


def build(spec, style):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    refs, cited = check_paper(spec)
    cfg, palette, fonts = style["paper"], style["palette"], style["typography"]
    order = cited + [key for key in refs if key not in cited]
    numbers = {key: i + 1 for i, key in enumerate(order)}

    def citation(match):
        key = match[1]
        if cfg["citation_style"] == "numeric":
            return f"[{numbers[key]}]"
        ref = refs[key]
        return f"({ref['authors']}, {ref['year']})"

    def resolve(text):
        return CITATION.sub(citation, text)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(cfg["width_mm"]), Mm(cfg["height_mm"])
    for edge, value in cfg["margins_mm"].items():
        setattr(section, f"{edge}_margin", Mm(value))
    for name, size, color in (("Normal", cfg["body_pt"], palette["text"]),
                              ("Title", cfg["title_pt"], palette["accent"]),
                              ("Heading 1", cfg["heading_pt"], palette["accent"])):
        doc_style = doc.styles[name]
        doc_style.font.name = fonts["latin"]
        doc_style.font.size = Pt(size)
        doc_style.font.color.rgb = RGBColor.from_string(color[1:])
        doc_style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), fonts["cjk"])
        doc_style.paragraph_format.line_spacing = cfg["line_spacing"]
        doc_style.paragraph_format.space_after = Pt(6)
        doc_style.paragraph_format.widow_control = True
    doc.add_paragraph(spec["title"], "Title")
    if spec.get("authors"):
        doc.add_paragraph(spec["authors"])
    if spec.get("abstract"):
        doc.add_heading("Abstract", level=1)
        doc.add_paragraph(resolve(spec["abstract"]))
    for part in spec["sections"]:
        doc.add_heading(part["heading"], level=1)
        for paragraph in part["paragraphs"]:
            doc.add_paragraph(resolve(paragraph))
    if refs:
        doc.add_heading("References", level=1)
        for key in order:
            ref = refs[key]
            prefix = f"[{numbers[key]}] " if cfg["citation_style"] == "numeric" else ""
            doc.add_paragraph(f"{prefix}{ref['authors']} ({ref['year']}). {ref['title']}. {ref['source']}.")
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue(), cited


def main(skill_dir):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("content", type=Path)
    p.add_argument("--style")
    p.add_argument("--output", required=True, type=Path)
    p.add_argument("--force", action="store_true")
    a = p.parse_args()
    if a.output.suffix.lower() != ".docx":
        raise ValueError("Output must end in .docx")
    style, source = load_style(skill_dir, a.style)
    preflight_outputs([a.output, manifest_path(a.output)], a.force, [source, a.content])
    content, cited = build(read_json(a.content), style)
    atomic_bytes(a.output, content)
    write_json(manifest_path(a.output), artifact_manifest(style, source, "python-docx", citations=cited,
               citation_style=style["paper"]["citation_style"],
               warnings=["References are user-supplied and not independently verified. Numeric/author-year formatting is basic, not a journal CSL implementation. Render in Word/LibreOffice before delivery."]))
    print(a.output)
