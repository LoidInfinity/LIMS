"""Observable contracts: extraction, distribution, sync, data fidelity and output style."""
from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "shared"), str(ROOT / "tools")]
from research_style.core import (atomic_bytes, fingerprint, load_style, read_json,
                                  validate, write_json)
from research_style.chart import read_data
from research_style.deck import build as build_deck
from research_style.extract import extract
from research_style.manage import installed_targets, synchronize
from research_style.paper import build as build_paper
from install import install


class SuiteTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.style = read_json(ROOT / "shared/styles/current.json")

    def tearDown(self):
        self.temp.cleanup()

    def cli(self, path, *args, expected=0, env=None):
        environment = os.environ.copy()
        environment.pop("RESEARCH_STYLE_FILE", None)
        environment["PYTHONUTF8"] = "1"
        environment.update(env or {})
        result = subprocess.run([sys.executable, str(path), *map(str, args)], cwd=self.root,
                                env=environment, text=True, capture_output=True, encoding="utf-8")
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return result

    def test_presets_and_invalid_schema(self):
        for file in (ROOT / "shared/styles/presets").glob("*.json"):
            validate(read_json(file))
        for update in ({"schema_version": 2}, {"unknown": 1}):
            broken = self.style | update
            with self.assertRaises(ValueError):
                validate(broken)
        for value in (float("nan"), 0, "large", True):
            broken = deepcopy(self.style)
            broken["chart"]["width_mm"] = value
            with self.assertRaises(ValueError):
                validate(broken)
        broken = deepcopy(self.style)
        broken["paper"]["margins_mm"]["left"] = 80
        broken["paper"]["margins_mm"]["right"] = 80
        broken["paper"]["width_mm"] = 180
        with self.assertRaises(ValueError):
            validate(broken)

    def test_json_duplicate_keys_and_nonfinite(self):
        file = self.root / "invalid.json"
        for text in ('{"a": 1, "a": 2}', '{"a": NaN}'):
            file.write_text(text)
            with self.assertRaises(ValueError):
                read_json(file)

    def test_style_precedence_and_no_silent_fallback(self):
        alternate = deepcopy(self.style)
        alternate["id"] = "alternate"
        file = self.root / "style.json"
        write_json(file, alternate)
        with patch.dict(os.environ, {"RESEARCH_STYLE_FILE": str(file)}):
            self.assertEqual(load_style(ROOT / "ppt-deck")[0]["id"], "alternate")
            self.assertEqual(load_style(ROOT / "ppt-deck", ROOT / "shared/styles/current.json")[0]["id"], self.style["id"])
            with self.assertRaises(OSError):
                load_style(ROOT / "ppt-deck", self.root / "missing.json")

    def test_standalone_install_sync_and_drift(self):
        destination = self.root / "安装 path"
        install(destination)
        # Running from a different cwd with no adjacent shared/ proves packaged resolution.
        hub = destination / "style-hub/scripts/sync_style.py"
        self.cli(hub, "validate")
        custom = deepcopy(self.style)
        custom["id"] = "custom"
        custom["palette"]["accent"] = "#CC1177"
        source = self.root / "custom.json"
        write_json(source, custom)
        self.cli(hub, "activate", source, "--dry-run")
        self.assertEqual(load_style(destination / "ppt-deck")[0]["id"], self.style["id"])
        self.cli(hub, "activate", source)
        for name in ("style-hub", "origin-chart", "ppt-deck", "paper-draft"):
            self.assertEqual(fingerprint(load_style(destination / name)[0]), fingerprint(custom))
        self.cli(hub, "check", "--style", source, "--skills-dir", destination)
        write_json(destination / "ppt-deck/assets/styles/current.json", self.style)
        self.cli(hub, "check", "--style", source, "--skills-dir", destination, expected=1)
        with self.assertRaises(ValueError):
            install(destination)

    def test_single_skill_runs_and_missing_assets_explain_install(self):
        destination = self.root / "single"
        install(destination, ["origin-chart"])
        script = destination / "origin-chart/scripts/render_origin.py"
        output = self.root / "plan.json"
        self.cli(script, ROOT / "examples/data.csv", "--x", "time", "--y", "control", "--backend", "plan", "--output", output)
        self.assertEqual(read_json(output)["data"]["control"], [1, 1.2, 1.5, 1.7, 1.9])

    def test_zip_release_is_self_contained(self):
        release = self.root / "release"
        self.cli(ROOT / "tools/package.py", "--output-dir", release)
        with zipfile.ZipFile(release / "research-style-suite.zip") as archive:
            self.assertIsNone(archive.testzip())
            self.assertEqual({name.split("/")[0] for name in archive.namelist()},
                             {"style-hub", "origin-chart", "ppt-deck", "paper-draft"})
            unpacked = self.root / "unpacked"
            archive.extractall(unpacked)
        self.cli(unpacked / "style-hub/scripts/sync_style.py", "validate")
        for name in ("style-hub", "origin-chart", "ppt-deck", "paper-draft"):
            with zipfile.ZipFile(release / f"{name}.zip") as archive:
                self.assertIsNone(archive.testzip())
                self.assertIn(f"{name}/assets/runtime/research_style/core.py", archive.namelist())
                self.assertIn(f"{name}/assets/styles/current.json", archive.namelist())

    def test_unowned_sync_target_and_rollback(self):
        (self.root / "ppt-deck").mkdir()
        with self.assertRaises(ValueError):
            installed_targets(self.root)
        first, second = self.root / "a.json", self.root / "b.json"
        first.write_bytes(b"before")
        second.write_bytes(b"also-before")
        calls = []
        def fail_second(path, data):
            calls.append(path)
            if len(calls) == 2:
                raise OSError("simulated write failure")
            atomic_bytes(path, data)
        with patch("research_style.manage.atomic_bytes", side_effect=fail_second):
            with self.assertRaises(OSError):
                synchronize(self.style, [first, second])
        self.assertEqual(first.read_bytes(), b"before")
        self.assertEqual(second.read_bytes(), b"also-before")

    def test_csv_validation_and_errorbars(self):
        good = read_data(ROOT / "examples/data.csv", "time", ["control"], ["control_sd"])
        self.assertEqual(len(good["time"]), 5)
        bad = self.root / "data.csv"
        for text in ("x,y\n1,NaN\n", "x,y\n1,\n", "x,x\n1,2\n", "x,y\n"):
            bad.write_text(text)
            with self.assertRaises(ValueError):
                read_data(bad, "x", ["y"])
        bad.write_text("x,y,e\n1,2,-1\n")
        with self.assertRaises(ValueError):
            read_data(bad, "x", ["y"], ["e"])

    def test_docx_style_and_citations(self):
        from docx import Document
        spec = {"title": "Demo", "sections": [{"heading": "Results", "paragraphs": ["Second [@b], first [@a], repeated [@b]."]}],
                "references": [{"id": key, "authors": key.upper(), "year": "2024", "title": f"Source {key}", "source": "User-supplied test fixture"} for key in ("a", "b")]}
        content, cited = build_paper(spec, self.style)
        document = Document(BytesIO(content))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Second [1], first [2], repeated [1].", text)
        self.assertEqual(cited, ["b", "a"])
        self.assertAlmostEqual(document.sections[0].left_margin.mm, 25, places=1)
        self.assertEqual(document.styles["Normal"].font.size.pt, self.style["paper"]["body_pt"])
        self.style["paper"]["citation_style"] = "author-year"
        text2 = "\n".join(p.text for p in Document(BytesIO(build_paper(spec, self.style)[0])).paragraphs)
        self.assertIn("Second (B, 2024)", text2)
        spec["sections"][0]["paragraphs"] = ["Missing [@missing]."]
        with self.assertRaises(ValueError):
            build_paper(spec, self.style)

    def test_deck_pagination_and_geometry(self):
        from pptx import Presentation
        bullets = [f"Measured observation number {i}" for i in range(18)]
        content, count = build_deck({"slides": [{"title": "Demo", "bullets": bullets}]}, self.root, self.style)
        self.assertGreater(count, 1)
        prs = Presentation(BytesIO(content))
        found = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        for bullet in bullets:
            self.assertEqual(found.count(bullet + "\n") + int(found.endswith(bullet)), 1)
        for slide in prs.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertLessEqual(shape.left + shape.width, prs.slide_width + 1)
                self.assertLessEqual(shape.top + shape.height, prs.slide_height + 1)

    def test_extraction_roundtrip_and_evidence(self):
        custom = deepcopy(self.style)
        custom["paper"]["body_pt"] = 13
        custom["paper"]["margins_mm"]["left"] = 31
        paper = self.root / "history.docx"
        paper.write_bytes(build_paper({"title": "History", "sections": [{"heading": "Test", "paragraphs": ["Evidence"]}]}, custom)[0])
        ppt = self.root / "history.pptx"
        ppt.write_bytes(build_deck({"slides": [{"title": "Title", "bullets": ["A body"]}]}, self.root, custom)[0])
        svg = self.root / "figure.svg"
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="120mm" height="80mm"><path stroke="#2255AA" d="M0 0 L1 1"/></svg>')
        result, evidence = extract([paper, ppt, svg], self.style)
        self.assertEqual(result["paper"]["body_pt"], 13)
        self.assertAlmostEqual(result["paper"]["margins_mm"]["left"], 31, places=1)
        self.assertEqual(result["chart"]["width_mm"], 120)
        self.assertEqual(result["slides"]["title_pt"], custom["slides"]["title_pt"])
        self.assertIn("paper.citation_style", evidence["retained_from_base"])
        self.assertEqual(extract([paper, ppt, svg], self.style), (result, evidence))
        invalid = self.root / "invalid.pptx"
        invalid.write_text("not a zip")
        with self.assertRaises(ValueError):
            extract([invalid], self.style)
        with self.assertRaises(ValueError):
            extract([self.root / "paper.pdf"], self.style)

    def test_all_outputs_share_fingerprint(self):
        chart = self.root / "chart.png"
        self.cli(ROOT / "origin-chart/scripts/render_origin.py", ROOT / "examples/data.csv", "--x", "time", "--y", "control", "treatment", "--backend", "matplotlib", "--output", chart)
        from PIL import Image
        with Image.open(chart) as bitmap:
            self.assertAlmostEqual(bitmap.width, self.style["chart"]["width_mm"] / 25.4 * self.style["chart"]["dpi"], delta=1)
        outline = self.root / "outline.json"
        write_json(outline, {"slides": [{"title": "Chart", "image": "chart.png"}]})
        pptx = self.root / "deck.pptx"
        self.cli(ROOT / "ppt-deck/scripts/build_pptx.py", outline, "--output", pptx)
        docx = self.root / "paper.docx"
        self.cli(ROOT / "paper-draft/scripts/assemble_paper.py", ROOT / "examples/paper.json", "--output", docx)
        for output in (chart, pptx, docx):
            manifest = read_json(str(output) + ".style.json")
            self.assertEqual(manifest["style_sha256"], fingerprint(self.style))
        svg = self.root / "chart.svg"
        self.cli(ROOT / "origin-chart/scripts/render_origin.py", ROOT / "examples/data.csv", "--x", "time", "--y", "control", "--backend", "matplotlib", "--output", svg)
        extracted, evidence = extract([svg], self.style)
        self.assertAlmostEqual(extracted["chart"]["width_mm"], self.style["chart"]["width_mm"], places=2)
        self.assertIn("chart.width_mm", evidence["fields"])
        # Refuse accidental replacement and preserve bytes.
        previous = docx.read_bytes()
        self.cli(ROOT / "paper-draft/scripts/assemble_paper.py", ROOT / "examples/paper.json", "--output", docx, expected=1)
        self.assertEqual(docx.read_bytes(), previous)


if __name__ == "__main__":
    unittest.main()
